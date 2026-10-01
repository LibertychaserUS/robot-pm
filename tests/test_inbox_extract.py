import json
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from robot_pm.inbox_extract import extract_file


def test_markdown_headings_become_sections(tmp_path: Path) -> None:
    path = tmp_path / "prd.md"
    path.write_text("# 发布\n\n正文在这里\n", encoding="utf-8")

    extracted = extract_file(path)

    assert extracted["kind"] == "text"
    assert extracted["sections"][0]["title"] == "发布"
    assert "正文在这里" in extracted["sections"][0]["body"]
    assert "发布" in extracted["text"]


def test_docx_heading_style_becomes_a_section(tmp_path: Path) -> None:
    path = tmp_path / "prd.docx"
    document = Document()
    document.add_heading("发布", level=1)
    document.add_paragraph("正文在这里")
    document.save(path)

    extracted = extract_file(path)

    assert extracted["kind"] == "text"
    assert extracted["sections"][0]["title"] == "发布"
    assert "正文在这里" in extracted["text"]


def test_pdf_without_text_layer_is_empty(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)

    extracted = extract_file(path)

    assert extracted == {"kind": "empty_text"}
    assert "text" not in extracted


def test_other_extension_is_unrecognized(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")

    assert extract_file(path) == {"kind": "unrecognized"}


def test_cli_prints_json(tmp_path: Path) -> None:
    path = tmp_path / "prd.md"
    path.write_text("# 一节\n\n内容\n", encoding="utf-8")
    from robot_pm.inbox_extract import main
    import io
    from contextlib import redirect_stdout

    buffer = io.StringIO()
    with redirect_stdout(buffer):
        code = main(["inbox_extract", str(path)])
    assert code == 0
    assert json.loads(buffer.getvalue())["kind"] == "text"
